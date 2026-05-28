from docx import Document

doc = Document()

doc.add_heading(
    "TERMO DE ABERTURA DE CONVÊNIO",
    level=1
)

doc.add_paragraph(
    """
GOVERNO DO ESTADO DO PIAUÍ
UNIVERSIDADE ESTADUAL DO PIAUÍ – UESPI
PRÓ-REITORIA DE EXTENSÃO, ASSUNTOS ESTUDANTIS E COMUNITÁRIOS – PREX
"""
)

doc.add_heading(
    "DADOS DA EMPRESA",
    level=2
)

doc.add_paragraph(
    """
Empresa: {{nome_empresa}}

CNPJ: {{cnpj}}

Endereço: {{endereco}}

Responsável legal: {{responsavel}}

Telefone: {{telefone}}

Tipo do convênio: {{tipo_convenio}}

Descrição: {{descricao}}
"""
)

doc.add_heading(
    "CLÁUSULA PRIMEIRA – DO OBJETIVO",
    level=2
)

doc.add_paragraph(
    """
O presente convênio tem como objetivo estabelecer
condições para viabilizar a concessão de estágio aos
discentes da UESPI.
"""
)

doc.add_heading(
    "CLÁUSULA SEGUNDA – FORMALIZAÇÃO",
    level=2
)

doc.add_paragraph(
    """
A formalização ocorrerá mediante assinatura do
Termo de Compromisso de Estágio.
"""
)

doc.save(
    "templates/termo_abertura_modelo.docx"
)

print("Modelo criado")
